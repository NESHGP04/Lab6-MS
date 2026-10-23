"""Construye el reporte a partir de resultados.json; requiere python-docx."""
from pathlib import Path
import json
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
data=json.loads((ROOT/'resultados.json').read_text())
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Heading 1','Heading 2']:
    s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
    s.font.size=Pt({'Normal':10.5,'Title':22,'Heading 1':16,'Heading 2':12}[name])
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.08

def p(text): return doc.add_paragraph(text)
def h(text): return doc.add_heading(text,level=1)
def sub(text): return doc.add_heading(text,level=2)
def page(text):
    para=h(text);para.paragraph_format.page_break_before=True
def fmt(v,n=4):
    if abs(v)<.5*10**(-n):v=0.
    return f'{v:,.{n}f}'
def interval(s,n=4): return f"[{fmt(s['low'],n)}, {fmt(s['high'],n)}]"
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1, cols=len(headers));t.autofit=False
    for i,x in enumerate(headers): t.rows[0].cells[i].text=str(x)
    for row in rows:
        cells=t.add_row().cells
        for c,x in zip(cells,row):c.text=str(x)
    for ri,row in enumerate(t.rows):
        for i,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[i])
            tcpr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            tcpr.append(borders)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DDE7EF' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF'));tcpr.append(shade)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2)
                for run in para.runs:run.font.size=Pt(9.5);run.bold=ri==0
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

doc.add_paragraph('Verificación del modelo híbrido de Aula Abierta',style='Title')
p('Task 2 completo  |  Grupo 4  |  q = 0.35 por día  |  D = 6 días')
p('Este reporte verifica por separado el clúster y la evolución de usuarios, compara la localización del evento de autoescalado y comprueba los conteos en la interfaz. Las verificaciones apoyan la implementación bajo los supuestos del enunciado; no determinan todavía una política óptima de costo y retención.')
h('Task 2.1 Verificación del clúster')
sub('a Comparación con M M c en régimen estable')
p('Se ejecutaron 30 réplicas independientes con λ = 6 solicitudes/s, μ = 2 solicitudes/s por servidor, c = 4 y K = 1 000 000. Cada réplica inicia vacía, dura 720 s y descarta 120 s. No se observaron rechazos. La capacidad elegida aproxima una cola sin límite.')
p('Con A = λ/μ = 3 y ρ = A/c = 0.75, la fórmula de Erlang C es:')
p('C(c,A) = [Aᶜ / (c! (1 − ρ))] / [Σ de k=0 a c−1 de Aᵏ/k! + Aᶜ / (c! (1 − ρ))].')
p('C(4,3) = 0.509434; Wq = C/(cμ − λ) = 0.254717 s; W = Wq + 1/μ = 0.754717 s. La utilización teórica es ρ = 0.75.')
stable=data['experiments']['stable'];th=data['theory_stable']
table(['Métrica','Teoría','Media simulada','IC del 95 %'],[[label,fmt(th[key]),fmt(stable['summary'][key]['mean']),interval(stable['summary'][key])] for label,key in [('Latencia W en s','W'),('Espera Wq en s','Wq'),('Utilización','util')]], [1.65,1.2,1.45,2.75])
p('Los intervalos usan la variabilidad entre réplicas: media ± t₀.₉₇₅,₂₉ × s/√30. No se tratan los tiempos de solicitudes correlacionadas de una misma cola como observaciones independientes.')
p('Los tres valores teóricos están dentro de sus intervalos. W y Wq quedan cerca del extremo inferior de sus intervalos, por lo que esta muestra es compatible con la teoría, pero no demuestra igualdad exacta ni ausencia de sesgo residual de calentamiento.')
p('W y Wq se calculan para la cohorte aceptada que llega entre los segundos 120 y 720. Tras 720 s se detienen las llegadas y se completa su atención. Este drenaje solo completa tiempos; no se añade a la duración medida, la utilización ni los conteos.')

page('Task 2.1 Sobrecarga y ley de Little')
sub('b Fracción rechazada con capacidad finita')
over=data['experiments']['overload']; ov=data['theory_overload']
p(f"Con λ = 10, c = 4 y K = 60, las 30 réplicas estiman r = {fmt(over['summary']['r']['mean'],6)}, con IC del 95 % {interval(over['summary']['r'],6)}. La aproximación de saturación es 1 − cμ/λ = 1 − 8/10 = 0.20.")
p('La aproximación supone que los cuatro servidores están ocupados todo el tiempo. Exactamente, λ(1 − r) = μ E[min(N,c)], de modo que r = 1 − μ E[min(N,c)]/λ. En una cola finita existe alguna probabilidad de tener servidores libres, por pequeña que sea.')
p(f"Como comprobación adicional, el proceso nacimiento y muerte M/M/4/60 da πₙ = πₙ₋₁ λ/[μ min(n,4)], normalizado sobre n = 0,…,60. Su resultado exacto estacionario es r = π₆₀ = {ov['r']:.10f}. Aquí la aproximación 0.20 resulta muy precisa. La utilización muestral fue 1 en todas las réplicas; eso no implica que su valor poblacional sea exactamente 1.")
sub('c Comprobación con la ley de Little')
p('Se miden L = ∫N(t)dt/600, Lq = ∫Q(t)dt/600 y la utilización ρ̂ = ∫B(t)dt/(600c). La tasa efectiva λ̂a es el número de llegadas aceptadas en la parte medida dividido entre 600. Las tres relaciones son L = λaW, Lq = λaWq y ρ = λa(W − Wq)/c.')
little=[]
for label,exp in [('λ = 6',stable),('λ = 10',over)]:
    rr=exp['rows']
    for name,left,prod in [('L','L',lambda r:r['lam_eff']*r['W']),('Lq','Lq',lambda r:r['lam_eff']*r['Wq']),('Utilización','util',lambda r:r['lam_eff']*(r['W']-r['Wq'])/4)]:
        measured=np.mean([r[left] for r in rr]); implied=np.mean([prod(r) for r in rr])
        little.append([label,name,fmt(measured,5),fmt(implied,5),f'{100*abs(measured-implied)/measured:.4f} %'])
table(['Caso','Magnitud','Medición','Por Little','Error relativo'],little,[.8,1.1,1.1,1.1,1.85])
p('Se comparan promedios de los productos por réplica, para conservar la covariación entre tasa y tiempo. El error relativo es |medición − producto|/medición. No se usa la tasa ofrecida cuando hay rechazos.')
p('Las pequeñas diferencias son efectos de borde: las áreas incluyen clientes presentes al segundo 120 y truncan al 720; la cohorte de llegadas observa tiempos completos. Además, el código verifica por separado la identidad exacta entre el área y la suma de las permanencias recortadas a [120,720]. Esta identidad debe cumplirse aun en una ventana finita y se cumple con error numérico menor que 10⁻⁷ segundos de área.')

page('Task 2.2 Verificación de usuarios')
sub('a RK4 propio frente a solve ivp')
p('Se fija δ = δ₀ = 0.003 por día y se usa U(0) = 100. La ecuación es f(U) = (0.002 + 0.35 U/20000)(20000 − U) − 0.003 U. La referencia es solve_ivp con DOP853, rtol = atol = 10⁻¹⁰ y salida continua.')
p('El RK4 implementado calcula k₁=f(U), k₂=f(U+hk₁/2), k₃=f(U+hk₂/2), k₄=f(U+hk₃) y U siguiente = U + h(k₁+2k₂+2k₃+k₄)/6. Euler calcula U siguiente = U + hf(U). El error máximo se toma en todos los nodos de cada malla durante 100 días.')
rows=data['sd']['rows']
table(['Paso h en días','Error máximo Euler','Error máximo RK4'],[[fmt(hh,4),f"{next(r['error'] for r in rows if r['h']==hh and r['method']=='Euler'):.7g}",f"{next(r['error'] for r in rows if r['h']==hh and r['method']=='RK4'):.7g}"] for hh in [1,.5,.25,.125,.0625]],[1.6,2.55,2.55])
sub('b Convergencia y selección del paso')
doc.add_picture(str(ROOT/'figuras/convergencia.png'),width=Inches(6.7))
sd=data['sd']
p(f"La pendiente de log(error) contra log(h), ajustada en los tres pasos menores, es {sd['orders']['Euler']:.3f} para Euler y {sd['orders']['RK4']:.3f} para RK4, próximas a los órdenes 1 y 4. La referencia entrega U(100) = {sd['U100']:.5f} usuarios.")
p(f"Se elige h = 0.25 días como paso fijo de referencia: RK4 tiene error máximo 0.005645 usuarios. En el híbrido se usa solve_ivp con las mismas tolerancias y max_step = 0.25 días, acortando el paso ante intercambios y entregas. Así se aísla el error de detección del evento. Como verificación independiente adicional, la solución cerrada de Riccati coincide con la referencia con error máximo {sd['reference_error']:.3g} usuarios.")

page('Task 2.3 Localización del autoescalado')
sub('a Eventos de estado y de tiempo')
p('El evento de estado es g(t,U) = aU/(cμ) − 0.7, con dirección positiva y detección habilitada solo si no hay un pedido pendiente. solve_ivp localiza la raíz dentro del paso y detiene la integración. El primer umbral con c = 4 equivale a U = 11 200 usuarios.')
p('Al cruzarlo se registra un pedido de dos servidores y se programa su entrega para t + 6 días. La integración termina exactamente en la entrega: c aumenta en dos y se elimina el pedido pendiente. Si aU/(cμ) sigue por encima de 0.7 con la nueva capacidad, se registra otro pedido de inmediato.')
p('Se intercambian datos con el DES en t = 0,1,…,99. W y r permanecen constantes hasta el intercambio siguiente, incluso si llegan servidores entre intercambios. Si coinciden eventos, primero se procesa la entrega y el reorden, después el intercambio y por último la revisión ingenua. Se cuentan las entregas hasta el día 100 inclusive; W máximo y r máximo corresponden a las ventanas que alimentan la evolución en [0,100).')
sub('b Comparación con revisión al final del paso')
p('Las versiones ingenuas revisan el umbral únicamente en la malla global de 1 o 0.25 días. Las tres versiones comparten integrador, H = 1, estado inicial y flujos aleatorios por ventana. La semilla base es 20261022. La tabla corresponde a una réplica emparejada.')
main=data['main_runs']
table(['Métrica','Evento localizado','Revisión 1 día','Revisión 0.25 días'],[[lab]+[fmt(r[k],5) if k!='c100' else str(r[k]) for r in main] for lab,k in [('Primer pedido en días','first_order'),('Primera entrega en días','first_delivery'),('Latencia máxima en s','Wmax'),('Usuarios en día 100','U100'),('Servidores en día 100','c100'),('Rechazo máximo','rmax')]],[1.9,1.65,1.65,1.65])
p('La misma secuencia base de tiempos entre llegadas y servicios se reutiliza por ventana entre versiones; las tasas y la capacidad pueden cambiar, de modo que los sistemas no tendrán necesariamente los mismos clientes aceptados. Esta técnica reduce parte del ruido de comparación sin eliminarlo.')
sub('c Diferencias y variabilidad')
table(['Versión menos localizada','Diferencia media','IC del 95 %'],[[label,fmt(data['paired'][mode][key]['mean'],5),interval(data['paired'][mode][key],5)] for mode,label,key in [('naive1','Pedido con revisión 1 día','first_order'),('naive025','Pedido con revisión 0.25 días','first_order'),('naive1','W máximo con revisión 1 día','Wmax'),('naive025','W máximo con revisión 0.25 días','Wmax'),('naive1','U(100) con revisión 1 día','U100'),('naive025','U(100) con revisión 0.25 días','U100')]],[2.8,1.3,2.65])
p('Estos intervalos se construyen con 30 diferencias emparejadas, no con una sola trayectoria. Las unidades de las diferencias son días, segundos o usuarios según la métrica.')

page('Task 2.3 Interpretación y límites')
doc.add_picture(str(ROOT/'figuras/eventos.png'),width=Inches(4.55))
a,b,c=main
p(f"En la réplica mostrada, la revisión diaria retrasa el primer pedido {b['first_order']-a['first_order']:.5f} días y la revisión cada 0.25 días lo retrasa {c['first_order']-a['first_order']:.5f} días. La entrega hereda ese retraso. Las diferencias de U(100) son {b['U100']-a['U100']:.3f} y {c['U100']-a['U100']:.3f} usuarios, respectivamente.")
p('En este experimento el efecto observado sobre W máximo y U(100) es nulo, salvo redondeo numérico, en las 30 comparaciones. La primera entrega exacta ocurre en 19.72057, la diaria en 20 y la de cuarto de día en 19.75: las tres ya están disponibles para el intercambio del día 20. Las segundas entregas también alimentan el mismo intercambio. Al mantener W y r constantes entre intercambios, H = 1 oculta el efecto del retraso. Los intervalos degenerados del primer pedido reflejan que antes del cruce todas las réplicas mantienen δ = δ₀; no representan incertidumbre nula en todo el modelo.')
p('La revisión tardía mantiene por más tiempo una capacidad insuficiente y puede elevar la latencia antes de la entrega. El daño es mayor si U crece rápidamente cerca del umbral y si D es grande respecto al tiempo de crecimiento: el sistema ya pasa varios días sin capacidad adicional y el retraso de detección se suma a ellos. En una ventana de congestión, δ puede superar temporalmente la tasa de adopción y reducir fuertemente U.')
p('No se debe exigir que W máximo o U(100) mejoren de manera estrictamente monótona al reducir el paso de revisión. El DES es aleatorio y los datos se actualizan una vez por día; mover una entrega respecto a ese calendario puede cambiar una ventana completa de W y r. Además, después de crecer la capacidad, las trayectorias de usuarios pueden acercarse nuevamente. Un U(100) similar no implica experiencias de servicio similares durante los 100 días.')
p('La verificación no demuestra que H = 1 sea inocuo ni que la política sea óptima. Cada ventana representa una hora pico bajo tasa constante, empieza vacía y pierde memoria entre llamadas. La ventana de 120 s de calentamiento reduce ese transitorio, pero no prueba que desaparezca en todos los regímenes. La sensibilidad a H y las mejoras respecto a la línea base pertenecen al análisis posterior.')

page('Task 2.4 Verificación de la interfaz')
sub('a Conteos de llegadas')
po=data['poisson']; interface=data['experiments']['interface']
p('Con λ = 5 se ejecutan 30 ventanas independientes. En la parte medida T = 600 s, el conteo debe tener distribución Poisson con media y varianza λT = 3000. Se cuentan todas las solicitudes generadas, incluidas las rechazadas, y no únicamente las atendidas.')
p(f"Media muestral = {fmt(po['mean'],3)}, con IC del 95 % {interval(interface['summary']['generated'],3)}; varianza muestral con divisor 29 = {fmt(po['variance'],3)}; cociente varianza/media = {fmt(po['dispersion'],4)}. Los valores esperados de media y varianza son ambos 3000.")
p(f"El estadístico aproximado de dispersión 29s²/media es {po['dispersion_stat']:.3f}; el rango central del 95 % de χ² con 29 grados de libertad es [{po['chi2_interval'][0]:.3f}, {po['chi2_interval'][1]:.3f}]. La media incluye 3000 en su intervalo; la varianza es menor que 3000 y la dispersión queda cerca del límite inferior, pero dentro del rango. Esta muestra no rechaza la compatibilidad con Poisson; 30 ventanas no demuestran la distribución.")
sub('b Conservación de entidades')
p('Para la parte medida, la identidad correcta es N inicial + llegadas = atendidas + rechazadas + N final. No puede omitirse N inicial porque al terminar el calentamiento hay solicitudes dentro. Para toda la ventana desde t = 0, el sistema empieza vacío y se cumple llegadas totales = atendidas totales + rechazadas totales + N final.')
balances=[]
for label,rr in [('5',interface['rows'][:5]),('10',over['rows'][:5])]:
    for r in rr: balances.append([f"{label} / {r['rep']}",r['initial'],r['generated'],r['departed'],r['rejected'],r['final'],r['balance']])
table(['λ / réplica','N inicial','Llegadas','Atendidas','Rechazos','N final','Balance'],balances,[1.,.75,1.,1.,.9,.75,.8])
p('El balance es cero en las 10 ventanas mostradas, incluidas cinco con sobrecarga, y en las 90 ventanas de las pruebas DES. El código comprueba ambas identidades en cada llamada, también en el híbrido. Las solicitudes completadas durante el drenaje no se cuentan como atendidas antes del segundo 720.')
sub('c Regla de calentamiento y sesgo')
p('Cada llamada comienza con el clúster vacío, utiliza la tasa y capacidad recibidas del SD y simula 720 s. Se descartan de las estadísticas de llegadas los primeros 120 s, pero no se eliminan los clientes presentes al comenzar la medición. Las áreas se calculan entre 120 y 720 s y los tiempos se atribuyen a las llegadas aceptadas en ese intervalo. Sin descarte, empezar vacío tiende a subestimar la espera, la latencia, la ocupación y el rechazo, especialmente en congestión. La regla reduce ese sesgo, pero no garantiza estacionariedad; para comprobarla haría falta variar el calentamiento y la duración.')
p('Reproducción: instalar numpy, scipy y matplotlib y ejecutar python task_2.py. Los CSV conservan las 30 réplicas de cada experimento, los errores numéricos, las comparaciones de eventos y las ventanas de la réplica principal. resultados.json conserva además las trayectorias y los intervalos. Este código corresponde al Task 2 y no sustituye el anexo sin modificar del Task 1.3.')

doc.core_properties.title='Verificación del modelo híbrido de Aula Abierta'
doc.core_properties.subject='Task 2 del grupo 4'
doc.core_properties.author=''
# Eliminar bordes heredados del estilo Title y de cualquier párrafo.
for root in [doc._element,doc.styles.element]:
    for node in list(root.iter(qn('w:pBdr'))):node.getparent().remove(node)
doc.save(ROOT/'Reporte_Task_2_Grupo_4.docx')
print('Reporte creado')
